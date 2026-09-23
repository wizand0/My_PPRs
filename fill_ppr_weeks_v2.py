#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fill_ppr_weeks_v2.py

Версия 2. Отличия от версии 1:
  1. Год из ячейки C1 листа "ППР" (например "2027 год") проставляется
     в ячейку E2 каждой техкарты.
  2. Определение начала раздела оборудования на листе "ППР" сделано
     надёжнее: разделом считается строка, где столбец A содержит ЦЕЛОЕ
     число (1, 2, 3... в т.ч. записанное текстом), а не признак "пусто
     в столбце C" (это ломалось на строках-продолжениях без адреса).
  3. Поиск ячейки "№ недели" на техкарте больше не требует, чтобы она
     была строго в одной строке с заголовком "ТО-N" — теперь ищется в
     нескольких строках после заголовка (это чинит случаи вроде
     "Вентиляция", где заголовок "ТО-1" и строка "№ недели" разнесены).
  4. Много больше предупреждений, по принципу "лучше лишнее
     предупреждение, чем тихая ошибка":
       - в ППР для раздела есть недели ТО-N, а на техкарте нет такого
         раздела (нет заголовка "ТО-N");
       - на техкарте есть раздел ТО-N, а в ППР для него не запланировано
         ни одной недели (блок останется пустым);
       - на техкарте есть заголовок "ТО-N", но рядом не удалось найти
         ячейку "№ недели" (раздел не будет заполнен);
       - раздел есть на листе "ППР", но НИ ОДНОГО листа-техкарты с таким
         наименованием в C3 не нашлось;
       - недель больше, чем свободных строк под разделом на техкарте;
       - совпадение ТО-2/ТО-3/ТО-4 на одной неделе (как в v1).

Использование:
    python3 fill_ppr_weeks_v2.py входной_файл.xlsx [выходной_файл.xlsx]

Как добавить новый раздел — без изменений от v1:
    1. Скопируйте существующий лист-техкарту как шаблон.
    2. В C3 впишите название раздела ТОЧНО как в столбце B на листе "ППР"
       (в строке-заголовке, где A — целое число).
    3. Запустите скрипт.
"""

import re
import sys
from pathlib import Path

import openpyxl

TO_NUM_RE = re.compile(r"^ТО-([1-4])\b")
WEEK_LABEL_KEYWORD = "недел"
BOUNDARY_KEYWORD = "перечень"  # "Перечень материалов при выполнении ТО:"
HEADER_NUM_RE = re.compile(r"^\d+$")  # целое число без точки -> строка-заголовок раздела
MAX_LABEL_SEARCH_ROWS = 6  # сколько строк максимум искать "№ недели" после заголовка ТО-N


def normalize(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def extract_year(ws):
    """Достаёт год (4 цифры) из ячейки C1 листа ППР."""
    v = ws["C1"].value
    if v is None:
        return None
    m = re.search(r"(\d{4})", str(v))
    return int(m.group(1)) if m else None


def build_ppr_map(ws):
    """
    Возвращает (result, original_names):
      result: dict normalized(group_name) -> {1:set(weeks), 2:.., 3:.., 4:..}
      original_names: normalized -> исходное (не нормализованное) название
    """
    max_row = ws.max_row
    max_col = ws.max_column

    # Найти строку с номерами недель (1..53).
    week_row = None
    data_first_col = None
    for r in range(1, min(10, max_row) + 1):
        for first_col_try in range(2, 6):
            cnt = 0
            total = 0
            for c in range(first_col_try, max_col + 1):
                v = ws.cell(row=r, column=c).value
                total += 1
                if isinstance(v, int) and 1 <= v <= 53:
                    cnt += 1
                elif isinstance(v, str) and v.strip().isdigit() and 1 <= int(v.strip()) <= 53:
                    cnt += 1
            if total > 0 and cnt / total > 0.6 and cnt >= 10:
                week_row = r
                data_first_col = first_col_try
                break
        if week_row:
            break

    if week_row is None:
        raise ValueError('На листе "ППР" не удалось найти строку с номерами недель (1..53).')

    data_start_row = week_row + 1

    # Определить принадлежность каждой строки к разделу.
    # Заголовок раздела = строка, где столбец A - целое число (без точки),
    # записанное числом ИЛИ текстом. Строки без такого A (в т.ч. полностью
    # пустой A, как бывает в строках-продолжениях) наследуют группу сверху.
    group_of_row = {}
    current_group = None
    for r in range(data_start_row, max_row + 1):
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        is_header = False
        if a is not None and HEADER_NUM_RE.match(str(a).strip()):
            is_header = True
        if is_header and b not in (None, ""):
            current_group = str(b).strip()
        group_of_row[r] = current_group

    result = {}
    original_names = {}
    to_values = {"то1": 1, "то2": 2, "то3": 3, "то4": 4}

    for r in range(data_start_row, max_row + 1):
        grp = group_of_row[r]
        if not grp:
            continue
        for c in range(data_first_col, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if not isinstance(v, str):
                continue
            key = v.strip().lower()
            if key in to_values:
                to_num = to_values[key]
                week_val = ws.cell(row=week_row, column=c).value
                try:
                    week_num = int(str(week_val).strip())
                except (TypeError, ValueError):
                    continue
                norm = normalize(grp)
                result.setdefault(norm, {1: set(), 2: set(), 3: set(), 4: set()})
                result[norm][to_num].add(week_num)
                original_names.setdefault(norm, grp)

    return result, original_names


def check_to_conflicts(ppr_map, original_names):
    """Недели, где для одного раздела совпадают 2+ типа из {ТО-2,ТО-3,ТО-4}."""
    conflicts = []
    for norm, data in ppr_map.items():
        all_weeks = set()
        for to_num in (2, 3, 4):
            all_weeks |= data.get(to_num, set())
        for week in sorted(all_weeks):
            present = [to_num for to_num in (2, 3, 4) if week in data.get(to_num, set())]
            if len(present) >= 2:
                to_list = ", ".join(f"ТО-{n}" for n in present)
                conflicts.append(
                    f'Раздел "{original_names[norm]}": на неделе {week} совпадают '
                    f"{to_list} — проверьте график на листе \"ППР\"."
                )
    return conflicts


def find_to_blocks(ws):
    """
    Возвращает (blocks, missing_label):
      blocks: dict to_num -> (week_col, fill_start_row, fill_end_row) включительно
      missing_label: список to_num, для которых заголовок "ТО-N" найден,
                     но рядом не нашлась ячейка "№ недели"
    """
    max_row = ws.max_row
    max_col = ws.max_column

    row_texts = {}
    headers = []  # (row, to_num)
    boundary_rows = []

    for r in range(1, max_row + 1):
        texts = {}
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                texts[c] = v
        if texts:
            row_texts[r] = texts
            for c, v in texts.items():
                m = TO_NUM_RE.match(v.strip())
                if m:
                    headers.append((r, int(m.group(1))))
                if BOUNDARY_KEYWORD in v.lower():
                    boundary_rows.append(r)

    if not headers:
        return {}, []

    headers.sort(key=lambda x: x[0])

    blocks = {}
    missing_label = []

    for i, (header_row, to_num) in enumerate(headers):
        next_header_row = headers[i + 1][0] if i + 1 < len(headers) else None
        search_upper = next_header_row if next_header_row else max_row + 1
        search_upper = min(search_upper, header_row + MAX_LABEL_SEARCH_ROWS)

        label_row = None
        week_col = None
        for r in range(header_row, min(search_upper, max_row) + 1):
            for c, v in row_texts.get(r, {}).items():
                if WEEK_LABEL_KEYWORD in v.lower():
                    label_row, week_col = r, c
                    break
            if label_row:
                break

        if not label_row:
            missing_label.append(to_num)
            continue

        # Если "№ недели" на той же строке, что и заголовок — данные начинаются
        # со следующей строки. Если ниже (уже на строке первого пункта работ) —
        # данные начинаются прямо там.
        start = label_row if label_row > header_row else label_row + 1

        if next_header_row:
            end = next_header_row - 1
        else:
            later_boundaries = [b for b in boundary_rows if b > header_row]
            end = (min(later_boundaries) - 1) if later_boundaries else max_row

        blocks[to_num] = (week_col, start, end)

    return blocks, missing_label


def fill_sheet(ws, ppr_map, original_names, warnings, year):
    if year is not None:
        ws["E2"].value = year

    group_cell = ws["C3"].value
    group_name = group_cell if group_cell not in (None, "") else ws.title
    norm = normalize(group_name)

    if norm not in ppr_map:
        warnings.append(
            f'Лист "{ws.title}": раздел "{group_name}" не найден на листе "ППР" '
            f"(проверьте точное совпадение названия в C3)."
        )
        return None

    blocks, missing_label = find_to_blocks(ws)

    for to_num in missing_label:
        warnings.append(
            f'Лист "{ws.title}": найден заголовок "ТО-{to_num}", но рядом (в пределах '
            f"{MAX_LABEL_SEARCH_ROWS} строк) не найдена ячейка \"№ недели\" — этот раздел "
            f"НЕ заполнен, проверьте структуру листа."
        )

    weeks_by_to = ppr_map[norm]

    for to_num in (1, 2, 3, 4):
        ppr_weeks = sorted(weeks_by_to.get(to_num, set()))

        if to_num in blocks:
            col, start, end = blocks[to_num]

            for r in range(start, end + 1):
                ws.cell(row=r, column=col).value = None

            capacity = end - start + 1
            weeks = ppr_weeks
            if len(weeks) > capacity:
                warnings.append(
                    f'Лист "{ws.title}", ТО-{to_num}: найдено {len(weeks)} недель, '
                    f"а свободно только {capacity} строк ({weeks[capacity:]} не поместились). "
                    f"Освободите больше строк под этим разделом."
                )
                weeks = weeks[:capacity]

            for i, week in enumerate(weeks):
                ws.cell(row=start + i, column=col).value = week

            if not ppr_weeks:
                warnings.append(
                    f'Лист "{ws.title}": есть раздел ТО-{to_num}, но в ППР для '
                    f'"{group_name}" не запланировано ни одной недели ТО-{to_num} — '
                    f"блок оставлен пустым (проверьте, так ли это должно быть)."
                )
        else:
            if to_num not in missing_label and ppr_weeks:
                # заголовка ТО-N на листе нет вообще (не просто не нашли label)
                warnings.append(
                    f'Лист "{ws.title}": в ППР для "{group_name}" есть ТО-{to_num} '
                    f'(недели: {", ".join(map(str, ppr_weeks))}), но на техкарте нет '
                    f"раздела ТО-{to_num} — добавьте его или проверьте структуру листа."
                )
            # если ppr_weeks пуст и блока нет вообще - это нормально
            # (например ТО-2/ТО-3 просто не предусмотрены для этого раздела)

    return norm


def main():
    if len(sys.argv) < 2:
        print("Использование: python3 fill_ppr_weeks_v2.py входной_файл.xlsx [выходной_файл.xlsx]")
        sys.exit(1)

    in_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else in_path

    wb = openpyxl.load_workbook(in_path, data_only=False)

    if "ППР" not in wb.sheetnames:
        print('Ошибка: в книге не найден лист "ППР".')
        sys.exit(1)

    ppr_ws = wb["ППР"]
    ppr_map, original_names = build_ppr_map(ppr_ws)
    year = extract_year(ppr_ws)

    conflicts = check_to_conflicts(ppr_map, original_names)

    warnings = []
    processed = []
    matched_norms = set()

    for sheet_name in wb.sheetnames:
        if sheet_name == "ППР":
            continue
        ws = wb[sheet_name]
        norm = fill_sheet(ws, ppr_map, original_names, warnings, year)
        if norm:
            matched_norms.add(norm)
        processed.append(sheet_name)

    unmatched = set(ppr_map.keys()) - matched_norms
    for norm in unmatched:
        warnings.append(
            f'Раздел "{original_names[norm]}" есть на листе "ППР", но не найден ни один '
            f"лист-техкарта с таким наименованием в C3 — данные по нему нигде не отображаются."
        )

    wb.save(out_path)

    print(f"Год из ППР!C1: {year if year is not None else 'не найден'}")
    print(f"Обработано листов: {len(processed)} -> {processed}")
    print(f"Разделы, найденные на ППР: {sorted(original_names.values())}")

    if conflicts:
        print("\nКОНФЛИКТЫ ГРАФИКА (ТО-2/ТО-3/ТО-4 в одну неделю):")
        for w in conflicts:
            print(" - " + w)
    else:
        print("\nКонфликтов ТО-2/ТО-3/ТО-4 по неделям не найдено.")

    if warnings:
        print(f"\nПРЕДУПРЕЖДЕНИЯ ({len(warnings)}):")
        for w in warnings:
            print(" - " + w)
    else:
        print("\nОстальное прошло без предупреждений.")

    print(f"\nСохранено: {out_path}")


if __name__ == "__main__":
    main()
